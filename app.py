"""
Flask application that provides the following services:
1) AI Resume Chatbot (using OpenAI GPT-3.5-Turbo)
2) CV Classification (using a pre-trained SVM model)
3) CV Analysis & Rating (using Google Generative AI / Gemini)

Notes:
- The application loads the necessary models and vectorizers for CV classification.
- It uses OpenAI's ChatCompletion for AI Resume Chatbot functionality.
- It uses Google Generative AI (Gemini) for CV analysis, rating, job recommendation, and improvement suggestions.
- Make sure to install all required packages (Flask, openai, joblib, docx, google.generativeai, etc.) before running.
"""

from flask import Flask, render_template, request, jsonify, send_file
import openai
import joblib
import os
from docx import Document
import google.generativeai as genai
from resume_extractor import extract_text_from_pdf, extract_resume_info
import re

# Configure Google Generative AI with your API key
genai.configure(api_key="AIzaSyC9z9vVcz6Nwo1O-i1EZsHkT571ZztS_i8")
model = genai.GenerativeModel(model_name="gemini-2.0-flash")

# Initialize Flask application
app = Flask(__name__)

# Upload folder configuration
UPLOAD_FOLDER = "uploads"
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

# Create 'uploads' folder if it doesn't exist
if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)

# Load pre-trained TF-IDF vectorizer and SVM model for CV classification
loaded_tfidf = joblib.load("services/tfidf_vectorizer.pkl")
loaded_svm = joblib.load("services/svm_model.pkl")

# Category mapping for classification results
category_mapping = {
    15: "Java Developer", 
    23: "Testing", 
    8: "DevOps Engineer", 
    20: "Python Developer",
    24: "Web Designing", 
    12: "HR", 
    13: "Hadoop", 
    3: "Blockchain", 
    10: "ETL Developer",
    18: "Operations Manager", 
    6: "Data Science", 
    22: "Sales", 
    16: "Mechanical Engineer",
    1: "Arts", 
    7: "Database", 
    11: "Electrical Engineering", 
    14: "Health and fitness",
    19: "PMO", 
    4: "Business Analyst", 
    9: "DotNet Developer", 
    2: "Automation Testing",
    17: "Network Security Engineer", 
    21: "SAP Developer", 
    5: "Civil Engineer", 
    0: "Advocate"
}

# Set OpenAI API key (ideally stored in an environment variable)
openai.api_key = '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'

# Global variable to store generated CV text from AI Resume Chatbot
cv_text = ""  

# -----------------------
#         Routes
# -----------------------

@app.route("/")
def home():
    """
    Render the main homepage (index.html).
    """
    return render_template("index.html")


@app.route("/ai-resume-chatbot")
def ai_resume_chatbot_page():
    """
    Render the AI Resume Chatbot page (GET request).
    """
    return render_template("ai-resume-chatbot.html")


@app.route("/ai-resume-chatbot", methods=["POST"])
def ai_resume_chatbot():
    """
    Handle POST requests for AI Resume Chatbot.
    - mode: "creation" or "inquiry" (default: "creation")
    - user_input: text entered by the user
    """
    global cv_text
    data = request.get_json()
    user_input = data.get("message", "").strip()
    mode = data.get("mode", "creation")  # Default mode is 'creation'
    
    # If no user input, return an error message in Arabic
    if not user_input:
        return jsonify({"message": "📢 عذرًا، لازم تكتب شيء عشان أقدر أساعدك! جرب تكتب بياناتك 😊"})
    
    # Adjust the system message based on the chosen mode
    if mode == "creation":
        system_message = (
            "أنت مساعد ذكي لإنشاء السير الذاتية. "
            "تحيي المستخدم بلهجة سعودية، لكن تكتب السيرة النهائية باللغة الإنجليزية وبصيغة احترافية متوافقة مع ATS. "
            "قسّم السيرة إلى الأقسام التالية: , Education, Work Experience, Skills, Certifications, Personal Information). "
            "إذا كانت المعلومات غير كافية، ضع مكان الأقسام الناقصة [More details needed]."
        )
    elif mode == "inquiry":
        system_message = (
            "أنت مساعد ذكي متخصص في الإجابة على استفسارات طريقة إنشاء السير الذاتية. "
            "قدم إجابات مفصلة وواضحة حول الخطوات، المتطلبات، وأفضل الممارسات لإنشاء سيرة ذاتية احترافية."
        )
    else:
        system_message = "أنت مساعد ذكي لإنشاء السير الذاتية."
    
    # Call OpenAI API (ChatCompletion)
    response = openai.ChatCompletion.create(
        model="gpt-3.5-turbo",
        messages=[
            {"role": "system", "content": system_message},
            {"role": "user", "content": user_input}
        ]
    )
    
    # Store the generated CV text in a global variable
    cv_text = response["choices"][0]["message"]["content"]
    
    # Return the generated text as JSON
    return jsonify({"message": cv_text})


@app.route("/download_cv")
def download_cv():
    """
    Allows the user to download the generated CV as a DOCX file.
    - If no CV text is available, returns an error message in JSON.
    """
    global cv_text

    if not cv_text:
        return jsonify({"message": "❌ لا توجد بيانات لإنشاء السيرة الذاتية، حاول إرسال بياناتك أولاً!"})

    # Create a DOCX document and add the CV text
    doc = Document()
    doc.add_heading("السيرة الذاتية", level=1)

    paragraphs = cv_text.split("\n")
    for para in paragraphs:
        doc.add_paragraph(para)

    # Save the file locally (CV.docx)
    file_path = "CV.docx"
    doc.save(file_path)

    # Send the file as an attachment
    return send_file(file_path, as_attachment=True)


@app.route("/cv-classification", methods=["GET", "POST"])
def cv_classification():
    """
    Classify the user's resume text into a specific job category using an SVM model.

    - If POST, read resume text from the form, vectorize it, and predict the category.
    - If GET, just render the classification page with no prediction.
    """
    prediction = None

    if request.method == "POST":
        resume_text = request.form["resumeText"] 


        resume_tfidf = loaded_tfidf.transform([resume_text])

        predicted_class = loaded_svm.predict(resume_tfidf)[0]
        prediction = category_mapping.get(predicted_class, "Unknown")

    return render_template("cv-classification.html", prediction=prediction)


@app.route("/cv-analysis-rating", methods=["GET", "POST"])
def cv_analysis_rating():
    """
    Analyze the uploaded CV (PDF) using the resume_extractor module and Gemini AI.
    - Extract text, parse resume info, limit skills to 6 technical ones, and generate:
      1) AI-based score (ai_score)
      2) Recommended job title (ai_job)
      3) 6 improvement suggestions (ai_recommendations)
    - Renders 'cv-analysis-rating.html' with the results.
    """
    resume_data = None
    ai_score = None
    ai_job = None
    ai_recommendations = None

    if request.method == "POST":
        # Check if the file key exists in the request
        if "cv_file" not in request.files:
            return jsonify({"error": "❌ يرجى رفع ملف السيرة الذاتية!"})

        file = request.files["cv_file"]

        # Check if a file was actually selected
        if file.filename == "":
            return jsonify({"error": "❌ لم يتم تحديد أي ملف!"})

        if file:
            # Save the file to the upload folder
            file_path = os.path.join(app.config["UPLOAD_FOLDER"], file.filename)
            file.save(file_path)

            # Extract text from the uploaded PDF
            with open(file_path, "rb") as pdf_file:
                text = extract_text_from_pdf(pdf_file)
                resume_data = extract_resume_info(text)
            
            # Simple list of tech-related keywords to filter skills
            tech_keywords = [
                "html", "css", "javascript", "typescript", "python", "java", "c++", "c#",
                "php", "sql", "mysql", "oracle", "docker", "kubernetes", "devops",
                "aws", "azure", "gcp", "linux", "unix", "shell", "git", "github",
                "machine learning", "deep learning", "data science", "hadoop", "spark",
                "etl", "power bi", "tableau", "react", "angular", "vue", "node",
                "django", "flask", "laravel", "spring", ".net", "dotnet", "kotlin",
                "swift", "r ", "matlab", "go ", "golang", "rust", "ruby", "rails",
                "scala", "perl", "bash", "networking", "cyber security", "ethical hacking"
            ]
            
            # Filter resume_data.skills to include only tech skills (up to 6)
            if resume_data and hasattr(resume_data, "skills"):
                technical_skills = []
                for skill in resume_data.skills:
                    skill_lower = skill.lower()
                    if any(kw in skill_lower for kw in tech_keywords):
                        technical_skills.append(skill)
                resume_data.skills = technical_skills[:6]

            # ---------------------
            # Gemini AI Prompts
            # ---------------------

            # 1) Generate a score (out of 100) with a brief explanation
            prompt_score = f"""
قيم السيرة الذاتية التالية من 100.
اعتبر العوامل التالية: المهارات (تأكد من صحتها التقنية)، الخبرة، التعليم، الشهادات، والملائمة للوظائف.
أعطِ نتيجة رقمية مع تفسير مختصر.
السيرة الذاتية:
{text}
الصيغة:
- التقييم: XX/100
- التفسير: [شرح مختصر]
            """
            ai_score = model.generate_content(prompt_score).text.strip()

            # 2) Suggest the most suitable job
            prompt_job = f"""
بناءً على التفاصيل في السيرة الذاتية أدناه، ما هو العمل الأنسب لهذا المرشح؟
أعطِ مسمى وظيفي واحد مع نسبة مئوية للتوافق.
الصيغة:
- العمل المناسب: [مسمى وظيفي]
- النسبة: XX%
السيرة الذاتية:
{text}
            """
            ai_job = model.generate_content(prompt_job).text.strip()

            # 3) Provide 6 improvement recommendations
            prompt_recommendations = f"""
قم بتحليل السيرة الذاتية أدناه وقدم 6 توصيات لتحسينها.
يجب أن تكون التوصيات باللغة العربية، كل توصية في سطر منفصل، بدون شرح إضافي.
السيرة الذاتية:
{text}
            """
            ai_recommendations = model.generate_content(prompt_recommendations).text.strip()

    # Render the result page with extracted data and AI outputs
    return render_template(
        "cv-analysis-rating.html",
        resume_data=resume_data,
        ai_score=ai_score,
        ai_job=ai_job,
        ai_recommendations=ai_recommendations
    )

# Run the Flask app
if __name__ == "__main__":
    app.run(debug=True)
